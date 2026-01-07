#!/usr/bin/env python3
"""
快速验证测试环境是否正确配置
"""

import sys
import os

def check_dependencies():
    """检查必要的依赖项"""
    print("检查依赖项...")
    required = {
        'pytest': 'pytest',
        'docx': 'python-docx',
    }

    missing = []
    for module, package in required.items():
        try:
            __import__(module)
            print(f"  ✓ {package}")
        except ImportError:
            print(f"  ✗ {package} (未安装)")
            missing.append(package)

    if missing:
        print(f"\n错误: 缺少以下依赖项: {', '.join(missing)}")
        print("请运行: pip install -r tests/requirements.txt")
        return False

    return True

def check_project_structure():
    """检查项目结构"""
    print("\n检查项目结构...")

    required_paths = [
        ('.claude/skills/ra-doc-assit/scripts', '脚本目录'),
        ('tests', '测试目录'),
        ('tests/test_extract_quality_standards.py', '提取测试'),
        ('tests/test_fill_quality_standards.py', '填充测试'),
    ]

    all_exist = True
    for path, description in required_paths:
        if os.path.exists(path):
            print(f"  ✓ {description}")
        else:
            print(f"  ✗ {description} (未找到)")
            all_exist = False

    return all_exist

def check_script_imports():
    """检查脚本是否可以导入"""
    print("\n检查脚本导入...")

    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    scripts_dir = os.path.join(project_root, '.claude', 'skills', 'ra-doc-assit', 'scripts')
    sys.path.insert(0, scripts_dir)

    try:
        import extract_quality_standards
        print("  ✓ extract_quality_standards.py")
    except ImportError as e:
        print(f"  ✗ extract_quality_standards.py: {e}")
        return False

    try:
        import fill_quality_standards
        print("  ✓ fill_quality_standards.py")
    except ImportError as e:
        print(f"  ✗ fill_quality_standards.py: {e}")
        return False

    return True

def run_simple_test():
    """运行一个简单的测试"""
    print("\n运行简单测试...")

    try:
        # 创建一个简单的测试
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.text = "测试"

        if cell.text == "测试":
            print("  ✓ 基本文档操作正常")
            return True
        else:
            print("  ✗ 文档操作异常")
            return False
    except Exception as e:
        print(f"  ✗ 测试失败: {e}")
        return False

def main():
    """主函数"""
    print("=" * 50)
    print("RA 文档技能 - 测试环境验证")
    print("=" * 50)
    print()

    checks = [
        check_dependencies(),
        check_project_structure(),
        check_script_imports(),
        run_simple_test(),
    ]

    print("\n" + "=" * 50)
    if all(checks):
        print("✓ 所有检查通过！测试环境已就绪。")
        print("\n运行测试:")
        print("  pytest                    # 运行所有测试")
        print("  pytest -v                 # 详细输出")
        print("  ./tests/run_tests.sh      # 使用脚本运行")
        print("=" * 50)
        return 0
    else:
        print("✗ 某些检查失败。请修复上述问题后再运行测试。")
        print("=" * 50)
        return 1

if __name__ == "__main__":
    sys.exit(main())
